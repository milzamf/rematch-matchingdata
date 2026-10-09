import openpyxl
import json
import os

path = r'D:\Ilzam\Code\Codex\Latsar\mfd_25_2.xlsx'
wb = openpyxl.load_workbook(path, read_only=True)
sheet = wb[wb.sheetnames[0]]

rows = sheet.iter_rows(values_only=True)
headers = next(rows)

provinces = {}
regencies = {}
districts = {}
total_villages = 0

for row in rows:
    if not row or row[3] is None or row[5] is None or row[7] is None or row[9] is None:
        continue
        
    kdprov = str(row[3]).strip().zfill(2)
    nmprov = str(row[4]).strip() if row[4] else ''
    kdkab = str(row[5]).strip().zfill(2)
    nmkab = str(row[6]).strip() if row[6] else ''
    kdkec = str(row[7]).strip().zfill(3)
    nmkec = str(row[8]).strip() if row[8] else ''
    kddesa = str(row[9]).strip().zfill(3)
    nmdesa = str(row[10]).strip() if row[10] else ''
    
    total_villages += 1
    
    if kdprov not in provinces:
        provinces[kdprov] = {
            'kd_prov': kdprov,
            'nm_prov': nmprov
        }
        
    idkab = kdprov + kdkab
    if idkab not in regencies:
        try:
            is_kota = int(kdkab) >= 71
        except ValueError:
            is_kota = 'KOTA' in nmkab.upper()
            
        tipe = 'KOTA' if is_kota else 'KABUPATEN'
        clean_nmkab = nmkab.replace('KABUPATEN ', '').replace('KOTA ', '').strip()
        full_name = f"{'KOTA' if is_kota else 'KAB.'} {clean_nmkab}"
        
        regencies[idkab] = {
            'id_kab': idkab,
            'kd_prov': kdprov,
            'kd_kab': kdkab,
            'nm_prov': nmprov,
            'nm_kab': clean_nmkab,
            'tipe': tipe,
            'full_name': full_name,
        }
        
    idkec = kdprov + kdkab + kdkec
    if idkec not in districts:
        districts[idkec] = {
            'id_kec': idkec,
            'id_kab': idkab,
            'kd_prov': kdprov,
            'kd_kab': kdkab,
            'kd_kec': kdkec,
            'nm_kec': nmkec
        }

print(f"Total Provinces: {len(provinces)}")
print(f"Total Regencies: {len(regencies)}")
print(f"Total Districts: {len(districts)}")
print(f"Total Villages: {total_villages}")

out_dir = r"D:\Ilzam\Code\Codex\Latsar\batchmatch\backend\storage\app\wilayah"
os.makedirs(out_dir, exist_ok=True)

with open(os.path.join(out_dir, 'provinces.json'), 'w', encoding='utf-8') as f:
    json.dump(provinces, f, ensure_ascii=False, indent=2)

with open(os.path.join(out_dir, 'regencies.json'), 'w', encoding='utf-8') as f:
    json.dump(regencies, f, ensure_ascii=False, indent=2)

print("Saved provinces.json and regencies.json to:", out_dir)

# Check 3515
print("3515 check:", regencies.get('3515'))
